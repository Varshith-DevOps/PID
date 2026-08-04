import asyncio
import os
from pathlib import Path

SUPPORTED_MIME_TYPES = {
    "application/pdf",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
}
SUPPORTED_EXTENSIONS = {".pdf", ".doc", ".docx", ".txt", ".xlsx", ".xls"}


class ResumeParseError(Exception):
    pass


def validate_resume_metadata(file_name: str, mime_type: str | None, size_bytes: int, max_size_mb: int) -> None:
    ext = Path(file_name).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS or (mime_type and mime_type.lower() not in SUPPORTED_MIME_TYPES):
        raise ResumeParseError("UNSUPPORTED_RESUME_FORMAT")
    if size_bytes > max_size_mb * 1024 * 1024:
        raise ResumeParseError("RESUME_TOO_LARGE")


def _extract_pdf_fast(path: str) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        return "\n".join(text_parts).strip()
    except Exception:
        return ""


def _extract_docx_fast(path: str) -> str:
    try:
        import docx
        doc = docx.Document(path)
        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text:
                text_parts.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text for cell in row.cells if cell.text]
                if row_text:
                    text_parts.append(" | ".join(row_text))
        return "\n".join(text_parts).strip()
    except Exception:
        return ""


def _extract_xlsx_fast(path: str) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            text_parts = []
            for sheet in wb.worksheets:
                text_parts.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(cell) for cell in row if cell is not None]
                    if row_vals:
                        text_parts.append(" | ".join(row_vals))
            return "\n".join(text_parts).strip()
        finally:
            wb.close()
    except Exception:
        return ""


async def parse_resume_file(path: str, timeout_seconds: int = 20) -> str:
    async def _parse() -> str:
        suffix = Path(path).suffix.lower()
        if suffix == ".txt":
            return Path(path).read_text(encoding="utf-8", errors="ignore").strip()

        extracted_text = ""
        if suffix == ".pdf":
            extracted_text = _extract_pdf_fast(path)
        elif suffix == ".docx":
            extracted_text = _extract_docx_fast(path)
        elif suffix == ".xlsx":
            extracted_text = _extract_xlsx_fast(path)

        # If fast extraction succeeded and returned substantial text (> 50 chars), return it immediately
        if len(extracted_text) > 50:
            return extracted_text

        # Otherwise, fall back to docling if installed (supports scanned documents/OCR)
        try:
            from docling.document_converter import DocumentConverter
            result = DocumentConverter().convert(path)
            return result.document.export_to_markdown()
        except ImportError:
            # If docling is not installed, return whatever fast extraction found, or raise error if empty
            if extracted_text:
                return extracted_text
            raise ResumeParseError("RESUME_PARSE_FAILED_NO_OCR_ENGINE")
        except Exception as exc:
            if extracted_text:
                return extracted_text
            raise ResumeParseError("RESUME_PARSE_FAILED") from exc

    if not os.path.exists(path):
        raise ResumeParseError("RESUME_NOT_FOUND")
    try:
        return await asyncio.wait_for(_parse(), timeout=timeout_seconds)
    except asyncio.TimeoutError as exc:
        raise ResumeParseError("RESUME_PARSE_FAILED") from exc
