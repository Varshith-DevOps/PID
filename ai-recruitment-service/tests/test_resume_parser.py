import os
import tempfile
import pytest
from unittest.mock import MagicMock, patch
from app.services.resume_parser import (
    parse_resume_file,
    validate_resume_metadata,
    ResumeParseError
)


def test_validate_resume_metadata():
    # Valid metadata should pass without raising error
    validate_resume_metadata("test.pdf", "application/pdf", 1024, 10)
    validate_resume_metadata("test.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", 1024, 10)
    validate_resume_metadata("test.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 1024, 10)
    validate_resume_metadata("test.txt", "text/plain", 1024, 10)

    # Invalid formats/sizes should raise ResumeParseError
    with pytest.raises(ResumeParseError) as exc:
        validate_resume_metadata("test.exe", "application/octet-stream", 1024, 10)
    assert str(exc.value) == "UNSUPPORTED_RESUME_FORMAT"

    with pytest.raises(ResumeParseError) as exc:
        validate_resume_metadata("test.pdf", "application/pdf", 20 * 1024 * 1024, 10)
    assert str(exc.value) == "RESUME_TOO_LARGE"


@pytest.mark.asyncio
async def test_parse_txt_file():
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as temp_file:
        temp_file.write("Hello, this is a plain text resume content.")
        temp_path = temp_file.name

    try:
        extracted = await parse_resume_file(temp_path)
        assert "Hello, this is a plain text resume content." in extracted
    finally:
        os.unlink(temp_path)


@pytest.mark.asyncio
async def test_parse_docx_file():
    import docx
    
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as temp_file:
        temp_path = temp_file.name
    
    doc = docx.Document()
    doc.add_paragraph("This is a Word document paragraph.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Header 1"
    table.rows[0].cells[1].text = "Value 1"
    doc.save(temp_path)

    try:
        extracted = await parse_resume_file(temp_path)
        assert "This is a Word document paragraph." in extracted
        assert "Header 1 | Value 1" in extracted
    finally:
        os.unlink(temp_path)


@pytest.mark.asyncio
async def test_parse_xlsx_file():
    import openpyxl
    
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as temp_file:
        temp_path = temp_file.name

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ResumeData"
    ws.cell(row=1, column=1, value="Name")
    ws.cell(row=1, column=2, value="John Doe")
    wb.save(temp_path)

    try:
        extracted = await parse_resume_file(temp_path)
        assert "--- Sheet: ResumeData ---" in extracted
        assert "Name | John Doe" in extracted
    finally:
        os.unlink(temp_path)


@pytest.mark.asyncio
async def test_parse_pdf_file_fast():
    # Mock pypdf PdfReader to return text to avoid writing valid PDF binary from scratch
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "This is text extracted from PDF page 1."
    
    mock_reader = MagicMock()
    mock_reader.pages = [mock_page]

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as temp_file:
        temp_path = temp_file.name

    try:
        with patch("pypdf.PdfReader", return_value=mock_reader):
            extracted = await parse_resume_file(temp_path)
            assert "This is text extracted from PDF page 1." in extracted
    finally:
        os.unlink(temp_path)


@pytest.mark.asyncio
async def test_parse_pdf_scanned_fallback_no_ocr():
    # PDF with no extractable text
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "" # Scanned/empty page
    
    mock_reader = MagicMock()
    mock_reader.pages = [mock_page]

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as temp_file:
        temp_path = temp_file.name

    try:
        with patch("pypdf.PdfReader", return_value=mock_reader):
            # If docling is not installed, it should raise or return empty.
            # We mock the import of docling to raise ImportError.
            with patch.dict("sys.modules", {"docling.document_converter": None}):
                with pytest.raises(ResumeParseError) as exc:
                    await parse_resume_file(temp_path)
                assert "RESUME_PARSE_FAILED_NO_OCR_ENGINE" in str(exc.value)
    finally:
        os.unlink(temp_path)
